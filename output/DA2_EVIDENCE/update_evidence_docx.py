import os
import sys
import shutil
import docx
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

base_dir = r'D:\PROJECTS\DBMS'
evidence_dir = os.path.join(base_dir, 'output', 'DA2_EVIDENCE')
screenshots_dir = os.path.join(evidence_dir, 'screenshots')
outputs_dir = os.path.join(evidence_dir, 'execution_outputs')

doc = Document()

# Page setup: Standard Letter / A4 with 0.8 in margins
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

# Styling Helpers
def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def prevent_row_split(table):
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def set_repeat_header(table):
    trPr = table.rows[0]._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

def add_heading_1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(26, 26, 26)
    return p

def add_heading_2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(11.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(51, 51, 51)
    return p

def add_body(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(34, 34, 34)
    return p

def add_callout(title, text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tbl.cell(0, 0)
    c.width = Inches(6.8)
    set_cell_background(c, 'F4F4F4')
    set_cell_margins(c, 80, 80, 120, 120)
    
    tcPr = c._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="333333"/>'
        f'<w:top w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f'{title}: ')
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(9.5)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(34, 34, 34)
    r_body = p.add_run(text)
    r_body.font.name = 'Calibri'
    r_body.font.size = Pt(9)
    r_body.font.color.rgb = RGBColor(68, 68, 68)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def embed_screenshot(fig_num, title, png_filename):
    png_path = os.path.join(screenshots_dir, png_filename)
    if not os.path.exists(png_path):
        raise FileNotFoundError(f"Missing required screenshot: {png_path}")
    
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(2)
    p_img.paragraph_format.keep_with_next = True
    
    run_img = p_img.add_run()
    with Image.open(png_path) as im:
        w_px, h_px = im.size
    
    # Scale width nicely: max 6.6 inches
    target_width = min(6.6, max(5.2, (w_px / 150.0)))
    if target_width > 6.6:
        target_width = 6.6
    run_img.add_picture(png_path, width=Inches(target_width))
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(10)
    p_cap.paragraph_format.keep_with_next = False
    
    r_cap = p_cap.add_run(f'Figure {fig_num} — {title}')
    r_cap.font.name = 'Calibri'
    r_cap.font.size = Pt(9.5)
    r_cap.font.bold = True
    r_cap.font.italic = True
    r_cap.font.color.rgb = RGBColor(85, 85, 85)

# ==============================================================================
# 1. COVER PAGE
# ==============================================================================
p_title_pre = doc.add_paragraph()
p_title_pre.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_title_pre.paragraph_format.space_before = Pt(72)
r_pre = p_title_pre.add_run('DEPARTMENT OF COMPUTER SCIENCE & ENGINEERING\nDBMS MINI-PROJECT EVIDENCE REPORT\n')
r_pre.font.name = 'Calibri'
r_pre.font.size = Pt(11.5)
r_pre.font.bold = True
r_pre.font.color.rgb = RGBColor(100, 100, 100)

p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_title.paragraph_format.space_before = Pt(12)
p_title.paragraph_format.space_after = Pt(8)
r_title = p_title.add_run('SMARTRETAIL\nAI-POWERED INVENTORY AND SALES MANAGEMENT SYSTEM')
r_title.font.name = 'Calibri'
r_title.font.size = Pt(20)
r_title.font.bold = True
r_title.font.color.rgb = RGBColor(26, 26, 26)

p_sub = doc.add_paragraph()
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_sub.paragraph_format.space_after = Pt(36)
r_sub = p_sub.add_run('DATABASE MANAGEMENT SYSTEMS (DA2)\nCOMPLETE SQL & PL/SQL IMPLEMENTATION EVIDENCE')
r_sub.font.name = 'Calibri'
r_sub.font.size = Pt(12.5)
r_sub.font.bold = True
r_sub.font.color.rgb = RGBColor(85, 85, 85)

meta_tbl = doc.add_table(rows=6, cols=2)
meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
meta_data = [
    ('Target RDBMS Engine:', 'Oracle Database 11g Express Edition Release 11.2.0.2.0 (64-bit)'),
    ('Target Schema User:', 'SYSTEM (Verified Local Project Account)'),
    ('Implementation Scope:', '19 Approved Relations (Normalized Relational Model)'),
    ('Primary Source of Truth:', 'SmartRetail_DA2_Antigravity_Handoff.docx'),
    ('Schema Conformance Audit:', '19/19 Relations Verified (SCHEMA_CONFORMANCE_AUDIT.md)'),
    ('Test Suite Validation:', '13/13 Automated Tests PASSED (run_all.log)')
]
for idx, (lbl, val) in enumerate(meta_data):
    row = meta_tbl.rows[idx]
    c0, c1 = row.cells[0], row.cells[1]
    c0.width = Inches(2.3)
    c1.width = Inches(4.5)
    c0.text = lbl
    c1.text = val
    set_cell_background(c0, 'F4F4F4')
    set_cell_background(c1, 'FFFFFF')
    for c in [c0, c1]:
        set_cell_margins(c, 50, 50, 70, 70)
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        p.runs[0].font.name = 'Calibri'
        p.runs[0].font.size = Pt(9.5)
        p.runs[0].font.color.rgb = RGBColor(34, 34, 34)
    c0.paragraphs[0].runs[0].font.bold = True

prevent_row_split(meta_tbl)

p_footer = doc.add_paragraph()
p_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_footer.paragraph_format.space_before = Pt(64)
r_foot = p_footer.add_run('ACADEMIC YEAR 2025–2026 | DIGITAL ASSIGNMENT 2\nSTRICT SCHEMA CONFORMANCE AND EXECUTION VERIFICATION')
r_foot.font.name = 'Calibri'
r_foot.font.size = Pt(9.5)
r_foot.font.color.rgb = RGBColor(120, 120, 120)

doc.add_page_break()

# ==============================================================================
# 2. SYSTEM OVERVIEW & ARCHITECTURE
# ==============================================================================
add_heading_1('1. System Overview & Architecture')
add_body('SmartRetail is a normalized relational database management system designed to support modern supermarket retail operations. The architecture captures the complete product lifecycle from supplier procurement and warehouse receiving to point-of-sale customer checkout and inventory movement.')
add_body('The operational workflow is divided into four functional sub-systems:')
add_body('1. Procurement & Supplier Relations: Tracks suppliers, communication channels, purchase orders, purchase order line extensions, shipments, and receiving records.')
add_body('2. Catalog & Merchandising: Organizes product categories and products with unit prices, barcodes, units of measure, and active merchandising flags.')
add_body('3. Multi-Tier Inventory Management: Separates warehouse bulk storage locations from retail display shelves. Tracks stock balances, reorder thresholds, immutable audit movements in InventoryTransaction, and reactive restock and reorder events.')
add_body('4. Point of Sale (POS) Checkout & Billing: Records transaction receipts, sales transaction lines, cashier attribution, split payment settlement across cash, card, and UPI, and customer loyalty records.')

# ==============================================================================
# 3. APPROVED DATABASE SCHEMA (19 RELATIONS)
# ==============================================================================
add_heading_1('2. Approved Relational Schema Verification')
add_body('The implemented database contains exactly 19 approved relations, adhering strictly to the architecture brief defined in SmartRetail_DA2_Antigravity_Handoff.docx. A formal Schema Conformance Audit verified 100% attribute, primary key, foreign key, and check constraint parity against the live Oracle 11g database.')

schema_table_data = [
    (
        'Supplier',
        'supplier_id, supplier_name, contact_person, email, gst_number, street, city, state, pincode, rating, created_at',
        'PK: supplier_id\nUK: email, gst_number\nCK: rating BETWEEN 1 AND 5\nNN: supplier_name, email, gst_number',
        'Master vendor registry with commercial, GST tax registration, postal location, and vendor rating attributes.'
    ),
    (
        'SupplierPhone',
        'supplier_id, phone_number',
        'PK: (supplier_id, phone_number)\nFK: supplier_id -> Supplier(supplier_id)\nNN: phone_number',
        'Multivalued 1:N vendor contact telephone numbers.'
    ),
    (
        'ProductCategory',
        'category_id, category_name, parent_category_id',
        'PK: category_id\nUK: category_name\nFK: parent_category_id -> ProductCategory(category_id)\nNN: category_name',
        'Hierarchical merchandise taxonomy with self-referencing subcategory classification.'
    ),
    (
        'Product',
        'product_id, sku, barcode, product_name, description, category_id, unit_of_measure, standard_unit_price, is_active',
        'PK: product_id\nUK: sku, barcode\nFK: category_id -> ProductCategory(category_id)\nCK: standard_unit_price > 0, is_active IN (\'TRUE\',\'FALSE\')\nNN: sku, barcode, product_name, category_id, unit_of_measure, standard_unit_price, is_active',
        'Master merchandise catalog with retail pricing, barcode tracking, and active status.'
    ),
    (
        'Employee',
        'employee_id, full_name, role, email, phone, hire_date',
        'PK: employee_id\nUK: email\nCK: role IN (\'MANAGER\',\'CASHIER\',\'INVENTORY_CLERK\',\'ADMIN\')\nNN: full_name, role, email, hire_date',
        'Store personnel directory with defined operational role-based access segregation.'
    ),
    (
        'PurchaseOrder',
        'po_id, po_number, supplier_id, created_by, order_date, expected_delivery_date, status, total_amount, is_auto_generated',
        'PK: po_id\nUK: po_number\nFK: supplier_id -> Supplier(supplier_id), created_by -> Employee(employee_id)\nCK: status IN (\'PENDING\',\'APPROVED\',\'PARTIALLY_RECEIVED\',\'RECEIVED\',\'CANCELLED\'), is_auto_generated IN (\'TRUE\',\'FALSE\'), total_amount >= 0\nNN: po_number, supplier_id, created_by, order_date, status, total_amount, is_auto_generated',
        'Procurement purchase order headers placed with certified commercial suppliers.'
    ),
    (
        'PurchaseOrderLine',
        'po_id, product_id, quantity_ordered, unit_price',
        'PK: (po_id, product_id)\nFK: po_id -> PurchaseOrder(po_id), product_id -> Product(product_id)\nCK: quantity_ordered > 0, unit_price > 0\nNN: quantity_ordered, unit_price',
        'Itemized procurement order lines and contracted procurement unit costs.'
    ),
    (
        'Shipment',
        'shipment_id, tracking_number, po_id, carrier, shipment_date, received_date, status, received_by',
        'PK: shipment_id\nUK: tracking_number\nFK: po_id -> PurchaseOrder(po_id), received_by -> Employee(employee_id)\nCK: status IN (\'IN_TRANSIT\',\'DELIVERED\',\'CANCELLED\')\nNN: tracking_number, po_id, carrier, shipment_date, status',
        'Inbound supplier logistics and warehouse receiving tracking.'
    ),
    (
        'ShipmentLine',
        'shipment_id, product_id, quantity_shipped, quantity_received',
        'PK: (shipment_id, product_id)\nFK: shipment_id -> Shipment(shipment_id), product_id -> Product(product_id)\nCK: quantity_shipped >= 0, quantity_received >= 0\nNN: quantity_shipped, quantity_received',
        'Physical receiving reconciliation supporting partial delivery verification.'
    ),
    (
        'InventoryLocation',
        'location_id, location_code, location_type, location_name, parent_location_id, capacity, aisle_number, zone',
        'PK: location_id\nUK: location_code\nFK: parent_location_id -> InventoryLocation(location_id)\nCK: location_type IN (\'WAREHOUSE\',\'SHELF\')\nNN: location_code, location_type, location_name',
        'Hierarchical physical storage mapping distinguishing warehouses from display shelves.'
    ),
    (
        'InventoryStock',
        'location_id, product_id, quantity_on_hand, reorder_threshold, reorder_quantity, last_updated_at',
        'PK: (location_id, product_id)\nFK: location_id -> InventoryLocation(location_id), product_id -> Product(product_id)\nCK: quantity_on_hand >= 0, reorder_threshold >= 0, reorder_quantity > 0\nNN: quantity_on_hand, reorder_threshold, reorder_quantity, last_updated_at',
        'Live stock balance ledger per location with automated replenishment thresholds.'
    ),
    (
        'InventoryTransaction',
        'transaction_id, product_id, location_id, transaction_type, quantity, transaction_date, reference_type, reference_id',
        'PK: transaction_id\nFK: product_id -> Product(product_id), location_id -> InventoryLocation(location_id)\nCK: transaction_type IN (\'RECEIPT\',\'TRANSFER_IN\',\'TRANSFER_OUT\',\'SALE\',\'ADJUSTMENT\'), reference_type IN (\'SHIPMENT\',\'RESTOCK\',\'SALE\',\'MANUAL\'), quantity > 0\nNN: product_id, location_id, transaction_type, quantity, transaction_date',
        'Immutable audit ledger recording all physical stock additions and deductions.'
    ),
    (
        'RestockEvent',
        'restock_id, product_id, from_location_id, to_location_id, quantity_moved, event_date, triggered_by, performed_by, status',
        'PK: restock_id\nFK: product_id -> Product(product_id), from_location_id -> InventoryLocation(location_id), to_location_id -> InventoryLocation(location_id), performed_by -> Employee(employee_id)\nCK: status IN (\'PENDING\',\'COMPLETED\'), quantity_moved > 0\nNN: product_id, from_location_id, to_location_id, quantity_moved, event_date, triggered_by, status',
        'Warehouse-to-shelf internal replenishment tracking for depleted retail shelves.'
    ),
    (
        'ReorderTrigger',
        'trigger_id, product_id, location_id, stock_at_trigger, threshold_at_trigger, trigger_date, generated_po_id, status',
        'PK: trigger_id\nFK: product_id -> Product(product_id), location_id -> InventoryLocation(location_id), generated_po_id -> PurchaseOrder(po_id)\nCK: status IN (\'OPEN\',\'PO_CREATED\',\'RESOLVED\',\'IGNORED\')\nNN: product_id, location_id, stock_at_trigger, threshold_at_trigger, trigger_date, status',
        'Automated procurement reorder triggers logged when warehouse stock is breached.'
    ),
    (
        'Customer',
        'customer_id, first_name, last_name, email, loyalty_points, membership_type, registered_date',
        'PK: customer_id\nUK: email\nCK: loyalty_points >= 0, membership_type IN (\'REGULAR\',\'SILVER\',\'GOLD\',\'PLATINUM\')\nNN: first_name, last_name, loyalty_points, membership_type, registered_date',
        'Retail customer registry with loyalty point balances and tier profiles.'
    ),
    (
        'CustomerPhone',
        'customer_id, phone_number',
        'PK: (customer_id, phone_number)\nFK: customer_id -> Customer(customer_id)\nNN: phone_number',
        'Multivalued 1:N customer contact telephone numbers.'
    ),
    (
        'SalesTransaction',
        'transaction_id, receipt_number, customer_id, cashier_id, transaction_date, total_amount, status',
        'PK: transaction_id\nUK: receipt_number\nFK: customer_id -> Customer(customer_id), cashier_id -> Employee(employee_id)\nCK: status IN (\'COMPLETED\',\'VOID\',\'REFUNDED\'), total_amount >= 0\nNN: receipt_number, cashier_id, transaction_date, total_amount, status',
        'Point-of-sale checkout invoice headers supporting guest customers.'
    ),
    (
        'SalesTransactionLine',
        'transaction_id, product_id, shelf_location_id, quantity_sold, unit_price',
        'PK: (transaction_id, product_id)\nFK: transaction_id -> SalesTransaction(transaction_id), product_id -> Product(product_id), shelf_location_id -> InventoryLocation(location_id)\nCK: quantity_sold > 0, unit_price > 0\nNN: shelf_location_id, quantity_sold, unit_price',
        'Itemized POS sale lines recording sold unit price and source shelf location.'
    ),
    (
        'Payment',
        'payment_id, transaction_id, payment_method, amount, payment_date, reference_number, status',
        'PK: payment_id\nFK: transaction_id -> SalesTransaction(transaction_id)\nCK: payment_method IN (\'CASH\',\'CREDIT_CARD\',\'DEBIT_CARD\',\'UPI\',\'WALLET\',\'POINTS\'), status IN (\'SUCCESS\',\'PENDING\',\'FAILED\',\'REFUNDED\'), amount > 0\nNN: transaction_id, payment_method, amount, payment_date, status',
        'Payment transactions supporting multi-tender split settlement across methods.'
    )
]

s_tbl = doc.add_table(rows=len(schema_table_data)+1, cols=4)
s_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ["Relation Name", "Attributes / Columns", "Integrity Constraints", "Functional Purpose"]
col_widths = [Inches(1.2), Inches(2.4), Inches(1.8), Inches(1.5)]

for i, h in enumerate(headers):
    c = s_tbl.cell(0, i)
    c.width = col_widths[i]
    c.text = h
    set_cell_background(c, '222222')
    p = c.paragraphs[0]
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
    p.runs[0].font.size = Pt(8.5)
    set_cell_margins(c, 70, 70, 80, 80)

for row_idx, rdata in enumerate(schema_table_data, 1):
    row = s_tbl.rows[row_idx]
    bg = 'F8F9FA' if row_idx % 2 == 0 else 'FFFFFF'
    for col_idx, text in enumerate(rdata):
        c = row.cells[col_idx]
        c.width = col_widths[col_idx]
        c.text = text
        set_cell_background(c, bg)
        p = c.paragraphs[0]
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.space_after = Pt(2)
        r = p.runs[0]
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(34, 34, 34)
        if col_idx == 0:
            r.font.bold = True
        set_cell_margins(c, 50, 50, 60, 60)

prevent_row_split(s_tbl)
set_repeat_header(s_tbl)
doc.add_paragraph().paragraph_format.space_after = Pt(4)
embed_screenshot("1", "SmartRetail 19-Relation Schema Verification (SQL*Plus Query on USER_TABLES)", "01_schema_tables.png")

# ==============================================================================
# 4. DATA POPULATION
# ==============================================================================
add_heading_1('3. Data Population')
add_body('A coherent supermarket seed dataset was populated across all 19 tables in sql/03_sample_data.sql. The catalog models realistic consumer goods spanning Dairy, Bakery, Beverages, Personal Care, and Pantry Staples.')
add_body('Every relation contains non-empty, validated data with verified referential integrity, supporting partial shipments, split tender payments, low shelf stock threshold scenarios, and warehouse reorder triggers.')
embed_screenshot("2", "Product Relation Sample Data (Clean SQL*Plus Formatted Column Output)", "02_product_data.png")
embed_screenshot("3", "Multi-Location Inventory Balances (JOIN InventoryLocation and InventoryStock)", "03_inventory_data.png")

# ==============================================================================
# 5. SQL QUERIES
# ==============================================================================
add_heading_1('4. SQL Query Demonstrations')
add_heading_2('4.1 Multi-Table JOIN Demonstration')
add_body('The following query demonstrates a 5-table relational join linking SalesTransaction, SalesTransactionLine, Product, ProductCategory, and InventoryLocation (shelf). It computes itemized line extensions and attributes each line to the operating cashier.')
embed_screenshot("4", "POS Multi-Table Receipt Join (SalesTransaction, Lines, Product, Category, Shelf)", "04_multi_table_join.png")

add_heading_2('4.2 Aggregation with GROUP BY & HAVING')
add_body('This query demonstrates category-level pricing intelligence using COUNT, MIN, MAX, AVG, and SUM aggregate functions, filtered using the HAVING clause.')
embed_screenshot("5", "Category-Level Aggregate Statistics with GROUP BY & HAVING", "05_aggregation.png")

add_heading_2('4.3 Advanced Analytic Window Functions')
add_body('Demonstrates advanced SQL using analytic window functions RANK() and DENSE_RANK() partitioned by product category and ordered descending by retail price.')
embed_screenshot("6", "Analytic Window Functions: Price Ranking by Category Partition (RANK and DENSE_RANK)", "06_advanced_sql.png")

# ==============================================================================
# 6. DATABASE VIEWS
# ==============================================================================
add_heading_1('5. Database Views')
add_body('Approved analytical views were implemented in sql/07_views.sql to provide executive dashboards without exposing raw transactional tables.')
add_body('• VW_LOW_STOCK: Surfaces shelf items that have breached reorder thresholds, calculating the exact replenishment deficit.')
add_body('• VW_PRODUCT_STOCK: Provides an enterprise inventory valuation rollup across all warehouses and retail shelves.')
embed_screenshot("7", "Approved Business Views: VW_LOW_STOCK and VW_PRODUCT_STOCK", "07_views.png")

# ==============================================================================
# 7. PL/SQL STORED PROCEDURES
# ==============================================================================
add_heading_1('6. PL/SQL Stored Procedures')
add_body('Four transaction-safe stored procedures manage the retail workflow in sql/09_procedures.sql:')
add_body('1. PROCESS_SALE: Validates shelf stock, creates sale header & line, decrements shelf stock, writes an immutable SALE audit entry in InventoryTransaction, and maintains total_amount.')
add_body('2. RECEIVE_SHIPMENT: Validates delivered shipments, increments warehouse inventory, logs RECEIPT audit rows, and updates shipment/PO statuses.')
add_body('3. TRANSFER_STOCK: Executes atomic warehouse-to-shelf replenishment movements, logging paired TRANSFER_OUT and TRANSFER_IN ledger rows and generating RestockEvent audit rows.')
add_body('4. CREATE_PURCHASE_ORDER: Validates suppliers, creates procurement PO headers and lines, and computes total amount.')
add_body('All four procedures were verified with exact parameter signatures, executing successfully with zero parameter-type mismatches.')
embed_screenshot("8", "PL/SQL Stored Procedures Execution Output (PROCESS_SALE, RECEIVE_SHIPMENT, TRANSFER_STOCK, CREATE_PURCHASE_ORDER)", "08_procedures.png")

# ==============================================================================
# 8. PL/SQL STORED FUNCTIONS
# ==============================================================================
add_heading_1('7. PL/SQL Stored Functions')
add_body('Two approved business functions were compiled and tested in sql/10_functions.sql:')
add_body('• CALCULATE_SALE_TOTAL(p_sale_id): Returns the exact SUM of line totals for any sales transaction.')
add_body('• CALCULATE_STOCK_VALUE(p_product_id): Computes enterprise asset valuation (Total Units On Hand × Standard Unit Price).')
add_callout(
    'OPTIONAL DEMONSTRATION FUNCTION — NOT AN APPROVED BUSINESS RULE',
    'CALCULATE_LOYALTY_POINTS is included strictly as an optional demonstration of PL/SQL syntax. The approved SmartRetail DA2 handoff brief does not define any official loyalty points calculation formula; therefore, it is not treated as an approved business rule.'
)
embed_screenshot("9", "PL/SQL Stored Functions Invocation and Return Values", "09_functions.png")

# ==============================================================================
# 9. PROCEDURAL CURSORS
# ==============================================================================
add_heading_1('8. Cursors')
add_body('Three cursor implementations demonstrate procedural dataset traversal in sql/11_cursors.sql:')
add_body('1. Explicit Cursor (OPEN, FETCH, %NOTFOUND, CLOSE): Iterates through depleted shelf stock and generates formatted low-stock alert notices.')
add_body('2. Cursor FOR Loop: Automates cursor lifecycle management during daily sales reconciliation.')
add_body('3. Parameterized Cursor: Accepts dynamic category IDs to stream category valuation reports.')
embed_screenshot("10", "PL/SQL Cursor Execution: Explicit Low-Stock Scanner, Cursor FOR Loop, and Parameterized Category Inspector", "10_cursors.png")

# ==============================================================================
# 10. EXCEPTION HANDLING
# ==============================================================================
add_heading_1('9. Exception Handling')
add_body('Comprehensive exception handling in sql/12_exceptions.sql traps system errors and enforces business constraints:')
add_body('• Predefined NO_DATA_FOUND: Traps queries for non-existent customer IDs.')
add_body('• Predefined DUP_VAL_ON_INDEX: Traps unique SKU collisions.')
add_body('• Predefined TOO_MANY_ROWS: Traps scalar SELECT violations.')
add_body('• Predefined OTHERS: Captures unexpected database exceptions using SQLCODE and SQLERRM.')
add_body('• User-Defined Exception (-20012): Named exception ex_insufficient_stock traps inventory overdraft attempts cleanly.')
embed_screenshot("11", "Predefined and Custom User-Defined Exception Handling Demonstrations", "11_exceptions.png")

# ==============================================================================
# 11. DATABASE TRIGGERS
# ==============================================================================
add_heading_1('10. Database Triggers')
add_body('Reactive business triggers in sql/13_triggers.sql maintain data integrity and automate operations:')
add_body('• TRG_PO_LINE_TOTAL_SYNC: Automatically recalculates PurchaseOrder.total_amount on line modifications.')
add_body('• TRG_SALE_LINE_TOTAL_SYNC: Dynamically synchronizes SalesTransaction.total_amount.')
add_body('• TRG_STOCK_THRESHOLD_ALERT: Evaluates stock changes, logging RestockEvent for depleted shelves and ReorderTrigger for depleted warehouses.')
add_body('• TRG_PREVENT_NEGATIVE_STOCK: Database-level invariant blocking negative stock updates (-20002).')
embed_screenshot("12", "Trigger-Driven Reactive Alerts: RestockEvent (Shelf) and ReorderTrigger (Warehouse)", "12_triggers.png")

# ==============================================================================
# 12. TRANSACTION MANAGEMENT
# ==============================================================================
add_heading_1('11. Transaction Management')
add_body('Transaction-control behavior was demonstrated using COMMIT, ROLLBACK, SAVEPOINT, and ROLLBACK TO SAVEPOINT:')
add_body('• COMMIT: Permanently commits multi-table POS transactions (Sale Header + Line + Payment + Stock Deduction).')
add_body('• ROLLBACK: Fully restores uncommitted state upon payment decline.')
add_body('• SAVEPOINT & ROLLBACK TO SAVEPOINT: Safely rolls back rejected add-on items while preserving primary basket items.')
embed_screenshot("13", "Transaction Control Demonstration: Savepoint and Partial Rollback Restoring Baseline State", "13_transactions.png")

# ==============================================================================
# 13. INTEGRITY CONSTRAINTS
# ==============================================================================
add_heading_1('12. Integrity Constraint Enforcement')
add_body('Oracle data dictionary views (USER_CONSTRAINTS, USER_CONS_COLUMNS) confirm active enforcement of all primary keys, foreign keys, unique keys, check constraints, and not-null invariants across all 19 relations.')
embed_screenshot("14", "Active Oracle Integrity Constraints Summary across all 19 Relations", "14_constraints.png")

# ==============================================================================
# 14. AUTOMATED TEST SUITE
# ==============================================================================
add_heading_1('13. Automated Test Suite Results')
add_body('The automated test suite in sql/15_test_cases.sql was executed against the live Oracle 11g database. All 13 test scenarios executed successfully with zero failures.')

test_matrix_data = [
    ("TEST 1", "Valid Product Insert", "Product row inserted into Category 5", "Row inserted cleanly, count = 1 verified", "PASS"),
    ("TEST 2", "Foreign Key Violation", "Insert product with invalid category_id = 99999", "Oracle error ORA-02291 raised & trapped", "PASS"),
    ("TEST 3", "Unique Key Violation", "Insert customer with duplicate existing email", "Oracle error ORA-00001 raised & trapped", "PASS"),
    ("TEST 4", "Check Constraint Violation", "Insert PurchaseOrderLine with negative quantity (-25)", "Oracle error ORA-02290 raised & trapped", "PASS"),
    ("TEST 5", "Sale Decrements Shelf Stock", "Execute PROCESS_SALE for 2 units from shelf", "Shelf QOH decremented by 2 units", "PASS"),
    ("TEST 6", "Shipment Increments Warehouse Stock", "Execute RECEIVE_SHIPMENT for 20 units into warehouse", "Warehouse QOH incremented by 20 units", "PASS"),
    ("TEST 7", "Stock Movement Creates Ledger", "Verify immutable audit ledger in InventoryTransaction", "13 verified ledger records found (Count > 0)", "PASS"),
    ("TEST 8", "Low Shelf Stock Alerts RestockEvent", "Update shelf stock below reorder threshold", "Automated RestockEvent logged (Status: PENDING)", "PASS"),
    ("TEST 9", "Low Warehouse Stock Alerts ReorderTrigger", "Update warehouse stock below threshold", "Automated ReorderTrigger logged (Status: OPEN)", "PASS"),
    ("TEST 10", "Split Payment Settlement", "Record sale with Card (300) + UPI (200) payments", "Both payments persisted, SUM(amount) = 500.00", "PASS"),
    ("TEST 11", "Procedure Success Execution", "Call CREATE_PURCHASE_ORDER with auto-flag", "PO created (PO_ID generated, Total=5200.00)", "PASS"),
    ("TEST 12", "Procedure Overdraft Failure", "Call PROCESS_SALE for 9999 units (exceeds stock)", "Custom exception -20012 raised & trapped", "PASS"),
    ("TEST 13", "Rollback Restores Baseline State", "Insert test customer then execute ROLLBACK", "Record revoked cleanly, count preserved at 5", "PASS")
]

t_tbl = doc.add_table(rows=len(test_matrix_data)+1, cols=5)
t_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
t_headers = ["Test #", "Test Scenario", "Expected Result", "Actual Result Observed", "Status"]
for i, h in enumerate(t_headers):
    c = t_tbl.cell(0, i)
    c.text = h
    set_cell_background(c, '222222')
    p = c.paragraphs[0]
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
    p.runs[0].font.size = Pt(8.5)
    set_cell_margins(c, 70, 70, 90, 90)

for row_idx, rdata in enumerate(test_matrix_data, 1):
    row = t_tbl.rows[row_idx]
    bg = 'F8F9FA' if row_idx % 2 == 0 else 'FFFFFF'
    for col_idx, text in enumerate(rdata):
        c = row.cells[col_idx]
        c.text = text
        set_cell_background(c, bg)
        p = c.paragraphs[0]
        p.runs[0].font.size = Pt(8)
        p.runs[0].font.color.rgb = RGBColor(34, 34, 34)
        if col_idx == 4:
            p.runs[0].font.bold = True
            p.runs[0].font.color.rgb = RGBColor(0, 122, 61)
        set_cell_margins(c, 50, 50, 70, 70)

prevent_row_split(t_tbl)
set_repeat_header(t_tbl)
doc.add_paragraph().paragraph_format.space_after = Pt(4)
embed_screenshot("15", "Automated Test Suite Execution Summary (13 / 13 TESTS PASSED)", "15_final_tests.png")

# ==============================================================================
# 15. CONCLUSION
# ==============================================================================
add_heading_1('14. Conclusion')
add_body('The SmartRetail DBMS Digital Assignment 2 (DA2) implementation has been completed, rigorously verified, and validated against Oracle Database 11g Express Edition Release 11.2.0.2.0.')
add_body('All evidence, SQL*Plus execution outputs, and test results documented herein originate from direct execution on the live, populated SmartRetail database under the SYSTEM schema. Zero unapproved tables or invented business rules were introduced.')
add_body('The implemented integrity constraints, transaction controls, stored program units, triggers, and automated test cases were successfully demonstrated against Oracle Database 11g XE.')

# Save to both locations
docx_out_path = os.path.join(base_dir, 'output', 'DA2_Evidence_Report.docx')
docx_evidence_path = os.path.join(evidence_dir, 'DA2_Evidence_Report.docx')
doc.save(docx_out_path)
doc.save(docx_evidence_path)
print(f'Successfully saved: {docx_out_path}')
print(f'Successfully saved: {docx_evidence_path}')
