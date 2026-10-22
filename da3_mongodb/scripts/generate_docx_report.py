"""
SmartRetail DA3 — Professional Academic DOCX Report Generator
Script: da3_mongodb/scripts/generate_docx_report.py
Description: Generates an academic submission report in Microsoft Word DOCX format
             matching the Markdown report content with professional styling.
"""

import os
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

base_dir = Path(__file__).resolve().parent.parent
md_report_path = base_dir / "reports" / "SmartRetail_DA3_Report.md"
docx_report_path = base_dir / "reports" / "SmartRetail_DA3_Report.docx"


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set cell margins in twips."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def set_cell_background(cell, fill_hex):
    """Set background color of a table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)


def build_docx_report():
    print(f"Reading Markdown source from: {md_report_path}")
    with open(md_report_path, "r", encoding="utf-8") as f:
        md_text = f.read()

    doc = Document()

    # Configure Margins (Normal 1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

    # Styling Palette
    PRIMARY_COLOR = RGBColor(26, 54, 93)     # Deep Navy
    SECONDARY_COLOR = RGBColor(43, 108, 176) # Slate Blue
    BODY_COLOR = RGBColor(45, 55, 72)        # Charcoal
    LIGHT_BG = "F7FAFC"
    BORDER_COLOR = "E2E8F0"

    # Title Page / Cover Section
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("SMARTRETAIL: AI-POWERED INVENTORY AND SALES MANAGEMENT SYSTEM")
    title_run.font.name = "Arial"
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = PRIMARY_COLOR
    title_p.paragraph_format.space_before = Pt(36)
    title_p.paragraph_format.space_after = Pt(12)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run("Digital Assignment 3 (DA3) — Academic Project Submission Report\nImplementation of SmartRetail Enterprise Database Using MongoDB")
    sub_run.font.name = "Arial"
    sub_run.font.size = Pt(14)
    sub_run.font.color.rgb = SECONDARY_COLOR
    sub_p.paragraph_format.space_after = Pt(30)

    # Metadata Box (Table)
    meta_table = doc.add_table(rows=6, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Course Title", "Database Management Systems (DBMS)"),
        ("Deliverable", "Digital Assignment 3 (DA3) — Non-Relational / Document DBMS"),
        ("Student Registration Number", "25BCE1074"),
        ("Student Name", "[Student Name]"),
        ("Academic Year", "2025–2026"),
        ("Faculty Guide", "[Faculty Name]")
    ]

    for idx, (label, val) in enumerate(meta_data):
        row = meta_table.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        
        lbl_p = cell_lbl.paragraphs[0]
        lbl_r = lbl_p.add_run(label)
        lbl_r.font.name = "Arial"
        lbl_r.font.bold = True
        lbl_r.font.size = Pt(10)
        lbl_r.font.color.rgb = PRIMARY_COLOR
        
        val_p = cell_val.paragraphs[0]
        val_r = val_p.add_run(val)
        val_r.font.name = "Arial"
        val_r.font.size = Pt(10)
        val_r.font.color.rgb = BODY_COLOR

        cell_lbl.width = Inches(2.5)
        cell_val.width = Inches(4.0)
        set_cell_margins(cell_lbl, 80, 80, 100, 100)
        set_cell_margins(cell_val, 80, 80, 100, 100)
        set_cell_background(cell_lbl, "EDF2F7")
        set_cell_background(cell_val, LIGHT_BG)

    doc.add_page_break()

    # Parse remaining markdown lines
    lines = md_text.split("\n")
    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    def flush_table(lines_to_flush):
        if not lines_to_flush:
            return
        # Parse rows
        parsed_rows = []
        for l in lines_to_flush:
            if "---" in l and "|" in l:
                continue  # separator row
            parts = [c.strip() for c in l.split("|")[1:-1]]
            if parts:
                parsed_rows.append(parts)

        if not parsed_rows:
            return

        num_cols = max(len(r) for r in parsed_rows)
        tbl = doc.add_table(rows=len(parsed_rows), cols=num_cols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

        for r_idx, row_data in enumerate(parsed_rows):
            is_header = (r_idx == 0)
            for c_idx in range(num_cols):
                cell = tbl.cell(r_idx, c_idx)
                text = row_data[c_idx] if c_idx < len(row_data) else ""
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                r = p.add_run(text)
                r.font.name = "Arial"
                r.font.size = Pt(9)
                if is_header:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(255, 255, 255)
                    set_cell_background(cell, "1A365D")
                else:
                    r.font.color.rgb = BODY_COLOR
                    bg_color = "F7FAFC" if r_idx % 2 == 1 else "FFFFFF"
                    set_cell_background(cell, bg_color)
                set_cell_margins(cell, 60, 60, 80, 80)
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    for line in lines[20:]:  # skip cover metadata already added
        stripped = line.strip()

        # Handle Code Blocks
        if stripped.startswith("```"):
            if in_code_block:
                # Flush code block
                code_text = "\n".join(code_lines)
                tbl = doc.add_table(rows=1, cols=1)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = tbl.cell(0, 0)
                cell.width = Inches(6.5)
                set_cell_background(cell, "2D3748")  # Dark code background
                set_cell_margins(cell, 100, 100, 120, 120)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                run = p.add_run(code_text)
                run.font.name = "Consolas"
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(226, 232, 240)  # Light gray text
                doc.add_paragraph().paragraph_format.space_after = Pt(6)
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        # Handle Tables
        if stripped.startswith("|") and stripped.endswith("|"):
            in_table = True
            table_lines.append(stripped)
            continue
        elif in_table:
            flush_table(table_lines)
            table_lines = []
            in_table = False

        if not stripped:
            continue

        # Headings
        if stripped.startswith("## "):
            h = doc.add_heading(stripped[3:], level=1)
            h.paragraph_format.space_before = Pt(16)
            h.paragraph_format.space_after = Pt(6)
            for r in h.runs:
                r.font.name = "Arial"
                r.font.color.rgb = PRIMARY_COLOR
                r.font.bold = True
        elif stripped.startswith("### "):
            h = doc.add_heading(stripped[4:], level=2)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
            for r in h.runs:
                r.font.name = "Arial"
                r.font.color.rgb = SECONDARY_COLOR
                r.font.bold = True
        elif stripped.startswith("#### "):
            h = doc.add_heading(stripped[5:], level=3)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
            for r in h.runs:
                r.font.name = "Arial"
                r.font.color.rgb = PRIMARY_COLOR
                r.font.bold = True
        elif stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(stripped[2:])
            r.font.name = "Arial"
            r.font.size = Pt(10)
            r.font.color.rgb = BODY_COLOR
        elif stripped.startswith("1. ") or stripped.startswith("2. ") or stripped.startswith("3. ") or stripped.startswith("4. ") or stripped.startswith("5. ") or stripped.startswith("6. ") or stripped.startswith("7. ") or stripped.startswith("8. "):
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(stripped[3:])
            r.font.name = "Arial"
            r.font.size = Pt(10)
            r.font.color.rgb = BODY_COLOR
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(stripped)
            r.font.name = "Arial"
            r.font.size = Pt(10)
            r.font.color.rgb = BODY_COLOR

    if in_table:
        flush_table(table_lines)

    # Save DOCX
    doc.save(docx_report_path)
    print(f"DOCX report successfully generated at: {docx_report_path}")


if __name__ == "__main__":
    build_docx_report()
